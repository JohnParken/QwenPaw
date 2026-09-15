"""Attempt-scoped Office and BI tools.

The functions in this module are deliberately boring: every file operation
goes through :class:`SkillExecutionContext`, scripts are argv-only subprocesses,
and publication is an explicit callback.  There is no general shell tool and
no path accepted from a caller is trusted before it is resolved.
"""
from __future__ import annotations

from dataclasses import dataclass, field
import csv
import hashlib
import json
import mimetypes
import os
from pathlib import Path
import selectors
import shutil
import signal
import subprocess
import sys
import time
from typing import Any, Callable, Iterable, Mapping, Sequence

from .bundle import BuiltinSkill, SkillManifest, discover_builtin_skills
from .verification import VerificationResult, verify_artifact


class ToolError(RuntimeError):
    """Base error for an attempt-scoped tool failure."""


class PathBoundaryError(ToolError):
    """Raised when a path is absolute, traverses, or leaves the attempt."""


class SkillScriptDenied(ToolError):
    """Raised when an argv command is not in the immutable manifest."""


class SkillScriptTimeout(ToolError):
    """Raised when a skill subprocess exceeds its execution deadline."""


class SkillOutputLimitExceeded(ToolError):
    """Raised when a subprocess writes more than its output budget."""


class SkillScriptError(ToolError):
    """Raised when an approved script exits unsuccessfully."""


@dataclass(frozen=True)
class ScriptResult:
    """Bounded result of a skill script invocation."""

    argv: tuple[str, ...]
    returncode: int
    stdout: str = ""
    stderr: str = ""
    duration_seconds: float = 0.0

    @property
    def output(self) -> str:
        return self.stdout + ("\n" if self.stdout and self.stderr else "") + self.stderr

    @property
    def ok(self) -> bool:
        return self.returncode == 0

    def as_dict(self) -> dict[str, Any]:
        return {
            "argv": list(self.argv),
            "returncode": self.returncode,
            "stdout": self.stdout,
            "stderr": self.stderr,
            "duration_seconds": self.duration_seconds,
        }


@dataclass(frozen=True)
class PublishedArtifact:
    """Metadata handed to the host callback after verification."""

    path: Path
    title: str
    mime_type: str
    size_bytes: int
    sha256: str
    verification: VerificationResult | None = None

    def as_dict(self) -> dict[str, Any]:
        return {
            "path": str(self.path),
            "title": self.title,
            "mime_type": self.mime_type,
            "size_bytes": self.size_bytes,
            "sha256": self.sha256,
            "verification": self.verification.as_dict() if self.verification else None,
        }


@dataclass(frozen=True)
class FileInfo:
    path: str
    size_bytes: int
    mime_type: str | None = None

    def as_dict(self) -> dict[str, Any]:
        return {
            "path": self.path,
            "size_bytes": self.size_bytes,
            "mime_type": self.mime_type,
        }


@dataclass(frozen=True)
class TableData:
    columns: tuple[str, ...]
    rows: tuple[tuple[Any, ...], ...]
    source: str = ""

    def as_dict(self) -> dict[str, Any]:
        return {
            "columns": list(self.columns),
            "rows": [list(row) for row in self.rows],
            "source": self.source,
        }


@dataclass(frozen=True)
class DocumentData:
    paragraphs: tuple[str, ...] = ()
    tables: tuple[tuple[tuple[str, ...], ...], ...] = ()
    source: str = ""

    def as_dict(self) -> dict[str, Any]:
        return {
            "paragraphs": list(self.paragraphs),
            "tables": [[list(row) for row in table] for table in self.tables],
            "source": self.source,
        }


@dataclass(frozen=True)
class PresentationData:
    slides: tuple[tuple[str, ...], ...] = ()
    source: str = ""

    def as_dict(self) -> dict[str, Any]:
        return {"slides": [list(slide) for slide in self.slides], "source": self.source}


@dataclass(frozen=True)
class PdfData:
    pages: tuple[str, ...] = ()
    source: str = ""

    def as_dict(self) -> dict[str, Any]:
        return {"pages": list(self.pages), "source": self.source}


def _safe_relative(value: str | Path) -> tuple[str, ...]:
    if isinstance(value, Path):
        raw = value.as_posix()
    elif isinstance(value, str):
        raw = value
    else:
        raise PathBoundaryError("path must be a string")
    if not raw or "\x00" in raw or "\\" in raw:
        raise PathBoundaryError("invalid path")
    path = Path(raw)
    if path.is_absolute() or ".." in path.parts:
        raise PathBoundaryError("path is outside attempt")
    parts = tuple(part for part in path.parts if part not in {"."})
    if not parts:
        raise PathBoundaryError("path is empty")
    return parts


def _check_components(root: Path, candidate: Path, *, allow_missing: bool) -> Path:
    try:
        root_real = root.resolve(strict=True)
        candidate_real = candidate.resolve(strict=not allow_missing)
        candidate_real.relative_to(root_real)
    except (OSError, ValueError) as exc:
        raise PathBoundaryError("path is outside attempt") from exc
    current = root
    try:
        relative = candidate.relative_to(root)
        for part in relative.parts:
            current = current / part
            if current.is_symlink():
                raise PathBoundaryError("symlink paths are not allowed")
    except ValueError as exc:
        raise PathBoundaryError("path is outside attempt") from exc
    return candidate_real


@dataclass(init=False)
class SkillExecutionContext:
    """The only filesystem and execution authority exposed to a Skill.

    ``root`` is an attempt directory.  It is created with the standard
    ``input``, ``work``/``workspace``, ``output`` and ``tmp`` children.  The
    immutable skill directory may live outside this root, but scripts always
    run with the attempt's work directory as their cwd.
    """

    root: Path
    skill_name: str
    skill_dir: Path | None
    work_dir: Path
    input_dir: Path
    output_dir: Path
    tmp_dir: Path
    manifest: SkillManifest | None
    timeout_seconds: float
    max_output_bytes: int
    max_file_bytes: int
    publish_callback: Callable[..., Any] | None

    def __init__(
        self,
        root: str | Path | None = None,
        skill: str | None = None,
        *,
        skill_name: str | None = None,
        skill_dir: str | Path | None = None,
        skill_root: str | Path | None = None,
        work_dir: str | Path | None = None,
        workspace_dir: str | Path | None = None,
        input_dir: str | Path | None = None,
        output_dir: str | Path | None = None,
        tmp_dir: str | Path | None = None,
        manifest: SkillManifest | Mapping[str, Any] | None = None,
        allowed_commands: Iterable[str] | None = None,
        entrypoints: Iterable[str] | None = None,
        timeout_seconds: float = 60.0,
        max_output_bytes: int = 65536,
        max_file_bytes: int = 256 * 1024 * 1024,
        publish_artifact: Callable[..., Any] | None = None,
        publish_callback: Callable[..., Any] | None = None,
    ) -> None:
        base = Path(root) if root is not None else None
        if base is None:
            candidate = work_dir or workspace_dir
            if candidate is None:
                raise PathBoundaryError("attempt root is required")
            base = Path(candidate).parent
        if base.exists() and base.is_symlink():
            raise PathBoundaryError("attempt root symlink is not allowed")
        base.mkdir(parents=True, exist_ok=True)
        self.root = base.resolve(strict=True)
        self.skill_name = skill_name or skill or ""
        raw_skill_dir = skill_dir or skill_root
        self.skill_dir = Path(raw_skill_dir).resolve(strict=True) if raw_skill_dir else None
        if self.skill_dir is not None and not self.skill_dir.is_dir():
            raise PathBoundaryError("skill directory is not a directory")
        raw_work = Path(work_dir or workspace_dir or (self.root / "work"))
        raw_input = Path(input_dir or (self.root / "input"))
        raw_output = Path(output_dir or (self.root / "output"))
        raw_tmp = Path(tmp_dir or (self.root / "tmp"))
        dirs = []
        for path in (raw_work, raw_input, raw_output, raw_tmp):
            resolved = _check_components(self.root, path, allow_missing=True)
            path.mkdir(parents=True, exist_ok=True)
            if path.is_symlink():
                raise PathBoundaryError("attempt directory symlink is not allowed")
            dirs.append(resolved)
        self.work_dir, self.input_dir, self.output_dir, self.tmp_dir = dirs
        if manifest is not None and not isinstance(manifest, SkillManifest):
            from .bundle import SkillManifest as _SkillManifest

            manifest = _SkillManifest.from_dict(manifest, skill_name=self.skill_name or None)
        self.manifest = manifest
        if manifest is not None:
            if allowed_commands is None:
                allowed_commands = manifest.allowed_commands
            if entrypoints is None:
                entrypoints = manifest.entrypoints
            timeout_seconds = min(timeout_seconds, manifest.timeout_seconds)
            max_output_bytes = min(max_output_bytes, manifest.max_output_bytes)
        self.allowed_commands = tuple(str(item) for item in (allowed_commands or ()))
        self.entrypoints = tuple(str(item) for item in (entrypoints or ()))
        if timeout_seconds <= 0 or max_output_bytes <= 0 or max_file_bytes <= 0:
            raise ToolError("execution limits must be positive")
        self.timeout_seconds = float(timeout_seconds)
        self.max_output_bytes = int(max_output_bytes)
        self.max_file_bytes = int(max_file_bytes)
        self.publish_callback = publish_artifact or publish_callback

    @property
    def workspace_dir(self) -> Path:
        return self.work_dir

    @property
    def allowed_roots(self) -> tuple[Path, ...]:
        return (self.input_dir, self.work_dir, self.output_dir, self.tmp_dir)

    def resolve_path(
        self,
        value: str | Path,
        *,
        writable: bool = False,
        area: str | None = None,
        must_exist: bool = False,
    ) -> Path:
        parts = _safe_relative(value)
        area_name = area or parts[0]
        aliases = {"workspace": "work", "work": "work"}
        area_name = aliases.get(area_name, area_name)
        roots = {
            "input": self.input_dir,
            "work": self.work_dir,
            "output": self.output_dir,
            "tmp": self.tmp_dir,
        }
        if area_name not in roots:
            raise PathBoundaryError("path must start in input, work, output, or tmp")
        if writable and area_name == "input":
            raise PathBoundaryError("input files are read-only")
        if area is None:
            base = self.root
            candidate = base.joinpath(*parts)
        else:
            base = roots[area_name]
            candidate = base.joinpath(*parts[1:] if parts[0] in roots or parts[0] == "workspace" else parts)
        return _check_components(self.root, candidate, allow_missing=not must_exist)

    def _resolve_existing_file(self, value: str | Path, *, area: str | None = None) -> Path:
        path = self.resolve_path(value, area=area, must_exist=True)
        if not path.is_file() or path.stat().st_size > self.max_file_bytes:
            raise ToolError("file is not readable")
        return path

    def list_files(self, directory: str = "input") -> list[FileInfo]:
        base = self.resolve_path(directory, must_exist=True)
        if not base.is_dir():
            raise ToolError("not a directory")
        result: list[FileInfo] = []
        for path in sorted(base.rglob("*")):
            if path.is_symlink():
                raise PathBoundaryError("symlink paths are not allowed")
            if not path.is_file():
                continue
            _check_components(self.root, path, allow_missing=False)
            relative = path.relative_to(self.root).as_posix()
            result.append(FileInfo(relative, path.stat().st_size, mimetypes.guess_type(path.name)[0]))
        return result

    def read_bytes(self, path: str | Path) -> bytes:
        return self._resolve_existing_file(path).read_bytes()

    def read_text(self, path: str | Path, *, encoding: str = "utf-8") -> str:
        value = self._resolve_existing_file(path)
        try:
            return value.read_text(encoding=encoding)
        except UnicodeDecodeError as exc:
            raise ToolError("file is not valid text") from exc

    def _prepare_write(self, path: str | Path) -> Path:
        target = self.resolve_path(path, writable=True)
        if target.exists() and target.is_dir():
            raise ToolError("destination is a directory")
        target.parent.mkdir(parents=True, exist_ok=True)
        _check_components(self.root, target.parent, allow_missing=False)
        return target

    def write_bytes(self, path: str | Path, content: bytes | bytearray) -> FileInfo:
        if not isinstance(content, (bytes, bytearray)):
            raise ToolError("binary content is required")
        if len(content) > self.max_file_bytes:
            raise ToolError("file exceeds size limit")
        target = self._prepare_write(path)
        target.write_bytes(bytes(content))
        return FileInfo(target.relative_to(self.root).as_posix(), len(content), mimetypes.guess_type(target.name)[0])

    def write_text(self, path: str | Path, content: str, *, encoding: str = "utf-8") -> FileInfo:
        if not isinstance(content, str):
            raise ToolError("text content is required")
        try:
            data = content.encode(encoding)
        except UnicodeEncodeError as exc:
            raise ToolError("content cannot be encoded") from exc
        return self.write_bytes(path, data)

    def copy_file(self, source: str | Path, destination: str | Path) -> FileInfo:
        source_path = self._resolve_existing_file(source)
        target = self._prepare_write(destination)
        shutil.copyfile(source_path, target)
        return FileInfo(target.relative_to(self.root).as_posix(), target.stat().st_size, mimetypes.guess_type(target.name)[0])

    def publish(
        self,
        path: str | Path,
        title: str | None = None,
        mime_type: str | None = None,
        *,
        verify: bool = True,
        requirements: dict[str, Any] | None = None,
    ) -> PublishedArtifact:
        target = self.resolve_path(path, must_exist=True)
        try:
            target.relative_to(self.output_dir)
        except ValueError as exc:
            raise PathBoundaryError("artifacts must be under output") from exc
        if target.is_symlink() or not target.is_file():
            raise ToolError("artifact is not a regular file")
        size = target.stat().st_size
        if size > self.max_file_bytes:
            raise ToolError("artifact exceeds size limit")
        if mime_type is None:
            mime_type = mimetypes.guess_type(target.name)[0] or "application/octet-stream"
        if requirements is None:
            requirements = {}
            if target.suffix.lower() in {".docx", ".pptx", ".pdf"}:
                requirements["render"] = True
            if target.suffix.lower() == ".xlsx":
                requirements.update({"render": True, "formula_recalculation": True})
        evidence = verify_artifact(target, requirements=requirements, expected_mime_type=mime_type) if verify else None
        if evidence is not None and not evidence.ok:
            raise ToolError("artifact verification failed: " + "; ".join(evidence.errors or evidence.limitations))
        digest = hashlib.sha256(target.read_bytes()).hexdigest()
        artifact = PublishedArtifact(target, title or target.name, mime_type, size, digest, evidence)
        if self.publish_callback is not None:
            callback = self.publish_callback
            try:
                callback_result = callback(target, artifact.title, artifact.mime_type)
            except TypeError:
                callback_result = callback(artifact)
            if callback_result is not None and isinstance(callback_result, PublishedArtifact):
                return callback_result
        return artifact


def _bundle_skill(context: SkillExecutionContext) -> BuiltinSkill | None:
    if not context.skill_name:
        return None
    try:
        return discover_builtin_skills().get(context.skill_name)
    except Exception:
        return None


def _canonical_command(value: str) -> str:
    name = Path(value).name.lower()
    if name.startswith("python3.") or name in {"python3", "python"}:
        return "python"
    if name.startswith("node") and name[4:].replace(".", "").isdigit():
        return "node"
    return name


_SHELL_COMMANDS = {"sh", "bash", "zsh", "fish", "dash", "csh", "ksh", "cmd", "powershell", "pwsh"}


def _entrypoint_from_argv(context: SkillExecutionContext, argv: Sequence[str]) -> tuple[int | None, str | None]:
    for index, token in enumerate(argv[1:], 1):
        if not isinstance(token, str) or "\x00" in token:
            raise SkillScriptDenied("invalid argv")
        if token in {"-c", "--eval", "-e"}:
            if context.entrypoints:
                raise SkillScriptDenied("INLINE_SCRIPT_NOT_ALLOWED")
            return None, None
        if token.startswith("-"):
            continue
        if "/" not in token and not token.endswith((".py", ".js", ".mjs", ".cjs")):
            continue
        if context.skill_dir is None:
            raise SkillScriptDenied("SKILL_DIRECTORY_REQUIRED")
        parts = _safe_relative(token)
        candidate = context.skill_dir.joinpath(*parts)
        try:
            resolved = _check_components(context.skill_dir, candidate, allow_missing=False)
        except PathBoundaryError as exc:
            raise SkillScriptDenied("ENTRYPOINT_OUTSIDE_SKILL") from exc
        if not resolved.is_file():
            raise SkillScriptDenied("ENTRYPOINT_NOT_FOUND")
        relative = resolved.relative_to(context.skill_dir).as_posix()
        if context.entrypoints and relative not in set(context.entrypoints):
            raise SkillScriptDenied("ENTRYPOINT_NOT_ALLOWED")
        if not context.entrypoints:
            raise SkillScriptDenied("ENTRYPOINT_NOT_REGISTERED")
        return index, relative
    return None, None


def _kill_process(process: subprocess.Popen[bytes]) -> None:
    try:
        if hasattr(os, "killpg"):
            os.killpg(process.pid, signal.SIGKILL)
        else:
            process.kill()
    except (OSError, ProcessLookupError):
        try:
            process.kill()
        except (OSError, ProcessLookupError):
            pass


def run_skill_script(
    context: SkillExecutionContext,
    argv: Sequence[str],
    *,
    timeout: float | None = None,
    timeout_seconds: float | None = None,
    max_output_bytes: int | None = None,
    output_limit: int | None = None,
    env: Mapping[str, str] | None = None,
) -> ScriptResult:
    """Run one manifest-approved command using ``shell=False``.

    The executable and entrypoint are checked before process creation.  The
    process group is killed on timeout or output overflow so a child cannot
    outlive the request.
    """
    if not isinstance(context, SkillExecutionContext):
        raise SkillScriptDenied("execution context is required")
    if isinstance(argv, (str, bytes, bytearray)) or not argv:
        raise SkillScriptDenied("argv must be a non-empty sequence")
    args = tuple(argv)
    if any(not isinstance(item, str) or not item or "\x00" in item for item in args):
        raise SkillScriptDenied("argv contains an invalid item")
    command = _canonical_command(args[0])
    if command in _SHELL_COMMANDS:
        raise SkillScriptDenied("SHELL_COMMAND_NOT_ALLOWED")
    allowed = {_canonical_command(item) for item in context.allowed_commands}
    if command not in allowed:
        raise SkillScriptDenied("COMMAND_NOT_ALLOWED")
    script_index, _ = _entrypoint_from_argv(context, args)
    if command == "python" and script_index is None and "-c" not in args and "-m" not in args:
        raise SkillScriptDenied("ENTRYPOINT_NOT_REGISTERED")
    actual = list(args)
    if command == "python" and Path(actual[0]).name.lower() in {"python", "python3"}:
        actual[0] = sys.executable
    if Path(actual[0]).is_absolute() and command != "python":
        resolved_executable = shutil.which(command)
        if resolved_executable is None or Path(actual[0]).resolve() != Path(resolved_executable).resolve():
            raise SkillScriptDenied("EXECUTABLE_NOT_ALLOWED")
    limit = max_output_bytes if max_output_bytes is not None else output_limit
    limit = context.max_output_bytes if limit is None else int(limit)
    if limit <= 0 or limit > context.max_output_bytes:
        raise SkillScriptDenied("OUTPUT_LIMIT_INVALID")
    requested_timeout = timeout_seconds if timeout_seconds is not None else timeout
    deadline = context.timeout_seconds if requested_timeout is None else float(requested_timeout)
    if deadline <= 0 or deadline > context.timeout_seconds:
        raise SkillScriptDenied("TIMEOUT_INVALID")
    child_env = {
        "PATH": os.environ.get("PATH", ""),
        "LANG": "C.UTF-8",
        "LC_ALL": "C.UTF-8",
        "PYTHONNOUSERSITE": "1",
        "PYTHONDONTWRITEBYTECODE": "1",
        "HOME": str(context.tmp_dir),
        "TMPDIR": str(context.tmp_dir),
    }
    if env:
        for key, value in env.items():
            if key not in {"PATH", "LANG", "LC_ALL", "PYTHONPATH", "PYTHONNOUSERSITE"} and not key.startswith("LC_"):
                raise SkillScriptDenied("ENVIRONMENT_NOT_ALLOWED")
            if not isinstance(value, str) or "\x00" in value:
                raise SkillScriptDenied("ENVIRONMENT_NOT_ALLOWED")
            child_env[key] = value
    start = time.monotonic()
    try:
        process = subprocess.Popen(
            actual,
            cwd=str(context.work_dir),
            env=child_env,
            stdin=subprocess.DEVNULL,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            shell=False,
            start_new_session=True,
        )
    except (OSError, ValueError) as exc:
        raise SkillScriptDenied(f"PROCESS_START_FAILED: {exc}") from exc

    selector = selectors.DefaultSelector()
    assert process.stdout is not None and process.stderr is not None
    selector.register(process.stdout, selectors.EVENT_READ, "stdout")
    selector.register(process.stderr, selectors.EVENT_READ, "stderr")
    chunks: dict[str, list[bytes]] = {"stdout": [], "stderr": []}
    total = 0
    timed_out = False
    exceeded = False
    try:
        while selector.get_map():
            remaining = deadline - (time.monotonic() - start)
            if remaining <= 0:
                timed_out = True
                _kill_process(process)
                break
            events = selector.select(min(remaining, 0.1))
            if not events and process.poll() is not None:
                # Give the pipes one last read before closing them.
                events = selector.select(0)
                if not events:
                    break
            for key, _ in events:
                try:
                    data = os.read(key.fileobj.fileno(), 8192)
                except OSError:
                    data = b""
                if not data:
                    selector.unregister(key.fileobj)
                    continue
                total += len(data)
                if total > limit:
                    exceeded = True
                    _kill_process(process)
                    break
                chunks[key.data].append(data)
            if timed_out or exceeded:
                break
    finally:
        if timed_out or exceeded:
            try:
                process.wait(timeout=1)
            except subprocess.TimeoutExpired:
                process.kill()
                process.wait()
        else:
            process.wait()
        selector.close()
    stdout = b"".join(chunks["stdout"]).decode("utf-8", "replace")
    stderr = b"".join(chunks["stderr"]).decode("utf-8", "replace")
    if timed_out:
        raise SkillScriptTimeout("SKILL_SCRIPT_TIMEOUT")
    if exceeded:
        raise SkillOutputLimitExceeded("SKILL_SCRIPT_OUTPUT_LIMIT")
    result = ScriptResult(tuple(actual), process.returncode or 0, stdout, stderr, time.monotonic() - start)
    if not result.ok:
        raise SkillScriptError(f"SKILL_SCRIPT_EXIT_{result.returncode}: {result.output[-4096:]}")
    return result


def list_files(context: SkillExecutionContext, directory: str = "input") -> list[FileInfo]:
    return context.list_files(directory)


def read_file(context: SkillExecutionContext, path: str | Path) -> bytes:
    return context.read_bytes(path)


def read_text(context: SkillExecutionContext, path: str | Path, *, encoding: str = "utf-8") -> str:
    return context.read_text(path, encoding=encoding)


def write_file(context: SkillExecutionContext, path: str | Path, content: bytes | bytearray) -> FileInfo:
    return context.write_bytes(path, content)


def write_text(context: SkillExecutionContext, path: str | Path, content: str, *, encoding: str = "utf-8") -> FileInfo:
    return context.write_text(path, content, encoding=encoding)


def copy_file(context: SkillExecutionContext, source: str | Path, destination: str | Path) -> FileInfo:
    return context.copy_file(source, destination)


def create_docx(
    context: SkillExecutionContext,
    path: str,
    title: str,
    paragraphs: Sequence[str],
) -> FileInfo:
    """Create a real DOCX in the output directory."""
    try:
        from docx import Document
    except ImportError as exc:
        raise ToolError("docx dependency is unavailable") from exc
    target = context._prepare_write(path)
    document = Document()
    if title:
        document.add_heading(title, level=0)
    for paragraph in paragraphs:
        document.add_paragraph(str(paragraph))
    document.save(str(target))
    return FileInfo(target.relative_to(context.root).as_posix(), target.stat().st_size, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")


def create_xlsx(
    context: SkillExecutionContext,
    path: str,
    columns: Sequence[str],
    rows: Sequence[Sequence[Any]],
    sheet_name: str = "Sheet1",
) -> FileInfo:
    """Create a real XLSX workbook from tabular values."""
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise ToolError("xlsx dependency is unavailable") from exc
    target = context._prepare_write(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = str(sheet_name)[:31] or "Sheet1"
    sheet.append([str(value) for value in columns])
    for row in rows:
        sheet.append(list(row))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(str(target))
    return FileInfo(target.relative_to(context.root).as_posix(), target.stat().st_size, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def create_pptx(
    context: SkillExecutionContext,
    path: str,
    title: str,
    slides: Sequence[Mapping[str, Any]],
) -> FileInfo:
    """Create a real PPTX using title-and-content layouts."""
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise ToolError("pptx dependency is unavailable") from exc
    target = context._prepare_write(path)
    presentation = Presentation()
    cover = presentation.slides.add_slide(presentation.slide_layouts[0])
    cover.shapes.title.text = str(title)
    for item in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = str(item.get("title", ""))
        body = slide.placeholders[1].text_frame
        body.clear()
        for index, bullet in enumerate(item.get("bullets", ())):
            paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
            paragraph.text = str(bullet)
    presentation.save(str(target))
    return FileInfo(target.relative_to(context.root).as_posix(), target.stat().st_size, "application/vnd.openxmlformats-officedocument.presentationml.presentation")


def create_pdf(
    context: SkillExecutionContext,
    path: str,
    title: str,
    paragraphs: Sequence[str],
) -> FileInfo:
    """Create a real text PDF with automatic page breaks."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
    except ImportError as exc:
        raise ToolError("PDF dependency is unavailable") from exc
    target = context._prepare_write(path)
    document = canvas.Canvas(str(target), pagesize=A4)
    width, height = A4
    y = height - 56
    document.setFont("Helvetica-Bold", 16)
    document.drawString(48, y, str(title)[:100])
    y -= 30
    document.setFont("Helvetica", 11)
    for paragraph in paragraphs:
        for line in str(paragraph).splitlines() or [""]:
            if y < 56:
                document.showPage()
                document.setFont("Helvetica", 11)
                y = height - 56
            document.drawString(48, y, line[:110])
            y -= 16
        y -= 6
    document.save()
    return FileInfo(target.relative_to(context.root).as_posix(), target.stat().st_size, "application/pdf")


def publish_artifact(
    context: SkillExecutionContext,
    path: str | Path,
    title: str | None = None,
    mime_type: str | None = None,
    *,
    verify: bool = True,
    requirements: dict[str, Any] | None = None,
) -> PublishedArtifact:
    return context.publish(path, title, mime_type, verify=verify, requirements=requirements)


def read_docx(context: SkillExecutionContext, path: str | Path) -> DocumentData:
    target = context._resolve_existing_file(path)
    try:
        from docx import Document

        document = Document(str(target))
        paragraphs = tuple(item.text for item in document.paragraphs)
        tables = tuple(
            tuple(tuple(cell.text for cell in row.cells) for row in table.rows)
            for table in document.tables
        )
        return DocumentData(paragraphs, tables, target.relative_to(context.root).as_posix())
    except ImportError as exc:
        raise ToolError("docx dependency is unavailable") from exc
    except Exception as exc:
        raise ToolError(f"DOCX_READ_FAILED: {exc}") from exc


def read_xlsx(context: SkillExecutionContext, path: str | Path) -> dict[str, TableData]:
    target = context._resolve_existing_file(path)
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(str(target), data_only=False, read_only=True)
        result: dict[str, TableData] = {}
        for sheet in workbook.worksheets:
            rows = [tuple(value for value in row) for row in sheet.iter_rows(values_only=True)]
            columns = tuple(str(value) if value is not None else "" for value in (rows[0] if rows else ()))
            result[sheet.title] = TableData(columns, tuple(rows[1:] if rows else ()), target.relative_to(context.root).as_posix())
        return result
    except ImportError as exc:
        raise ToolError("xlsx dependency is unavailable") from exc
    except Exception as exc:
        raise ToolError(f"XLSX_READ_FAILED: {exc}") from exc


def read_pptx(context: SkillExecutionContext, path: str | Path) -> PresentationData:
    target = context._resolve_existing_file(path)
    try:
        from pptx import Presentation

        presentation = Presentation(str(target))
        slides = tuple(
            tuple(shape.text for shape in slide.shapes if hasattr(shape, "text"))
            for slide in presentation.slides
        )
        return PresentationData(slides, target.relative_to(context.root).as_posix())
    except ImportError as exc:
        raise ToolError("pptx dependency is unavailable") from exc
    except Exception as exc:
        raise ToolError(f"PPTX_READ_FAILED: {exc}") from exc


def read_pdf(context: SkillExecutionContext, path: str | Path) -> PdfData:
    target = context._resolve_existing_file(path)
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(target), strict=False)
        pages = tuple(page.extract_text() or "" for page in reader.pages)
        return PdfData(pages, target.relative_to(context.root).as_posix())
    except ImportError as exc:
        raise ToolError("pdf dependency is unavailable") from exc
    except Exception as exc:
        raise ToolError(f"PDF_READ_FAILED: {exc}") from exc


def read_table(context: SkillExecutionContext, path: str | Path, *, delimiter: str | None = None) -> TableData:
    target = context._resolve_existing_file(path)
    suffix = target.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        values = read_xlsx(context, path)
        return next(iter(values.values()), TableData((), (), target.relative_to(context.root).as_posix()))
    if suffix == ".json":
        try:
            value = json.loads(target.read_text(encoding="utf-8"))
        except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ToolError("JSON_READ_FAILED") from exc
        if isinstance(value, Mapping):
            value = [value]
        if not isinstance(value, list) or not all(isinstance(row, Mapping) for row in value):
            raise ToolError("JSON table must be an array of objects")
        columns = tuple(dict.fromkeys(key for row in value for key in row))
        rows = tuple(tuple(row.get(column) for column in columns) for row in value)
        return TableData(columns, rows, target.relative_to(context.root).as_posix())
    if suffix == ".parquet":
        try:
            import pandas as pd

            frame = pd.read_parquet(target)
            columns = tuple(str(column) for column in frame.columns)
            rows = tuple(
                tuple(
                    None if pd.isna(value) else value.item() if hasattr(value, "item") else value
                    for value in row
                )
                for row in frame.itertuples(index=False, name=None)
            )
            return TableData(columns, rows, target.relative_to(context.root).as_posix())
        except ImportError as exc:
            raise ToolError("Parquet support requires pandas and pyarrow") from exc
        except Exception as exc:
            raise ToolError(f"PARQUET_READ_FAILED: {exc}") from exc
    sep = delimiter or ("\t" if suffix == ".tsv" else ",")
    try:
        with target.open(newline="", encoding="utf-8-sig") as stream:
            rows = [tuple(row) for row in csv.reader(stream, delimiter=sep)]
    except (OSError, UnicodeDecodeError, csv.Error) as exc:
        raise ToolError("TABLE_READ_FAILED") from exc
    columns = rows[0] if rows else ()
    return TableData(tuple(columns), tuple(rows[1:]), target.relative_to(context.root).as_posix())


def profile_table(context: SkillExecutionContext, path: str | Path) -> dict[str, Any]:
    table = read_table(context, path)
    missing = {
        column: sum(value in {None, ""} for value in (row[index] for row in table.rows))
        for index, column in enumerate(table.columns)
    }
    return {
        "source": table.source,
        "rows": len(table.rows),
        "columns": list(table.columns),
        "missing": missing,
        "filters": [],
    }


def aggregate_table(
    context: SkillExecutionContext,
    path: str,
    group_by: str,
    value_column: str,
    operation: str = "sum",
) -> dict[str, float]:
    """Aggregate an uploaded table without arbitrary code execution."""
    if operation not in {"sum", "mean", "count", "min", "max"}:
        raise ToolError("unsupported aggregation")
    table = read_table(context, path)
    try:
        group_index = table.columns.index(group_by)
        value_index = table.columns.index(value_column)
    except ValueError as exc:
        raise ToolError("aggregation column does not exist") from exc
    grouped: dict[str, list[float]] = {}
    for row in table.rows:
        try:
            number = float(row[value_index])
        except (TypeError, ValueError):
            continue
        grouped.setdefault(str(row[group_index]), []).append(number)
    result: dict[str, float] = {}
    for key, values in grouped.items():
        if operation == "sum":
            result[key] = sum(values)
        elif operation == "mean":
            result[key] = sum(values) / len(values)
        elif operation == "count":
            result[key] = float(len(values))
        elif operation == "min":
            result[key] = min(values)
        else:
            result[key] = max(values)
    return result


def pivot_table(
    context: SkillExecutionContext,
    path: str,
    row_column: str,
    column_column: str,
    value_column: str,
    operation: str = "sum",
) -> dict[str, dict[str, float]]:
    """Create a bounded two-dimensional pivot from uploaded tabular data."""
    if operation not in {"sum", "mean", "count", "min", "max"}:
        raise ToolError("unsupported pivot aggregation")
    table = read_table(context, path)
    try:
        row_index = table.columns.index(row_column)
        column_index = table.columns.index(column_column)
        value_index = table.columns.index(value_column)
    except ValueError as exc:
        raise ToolError("pivot column does not exist") from exc
    grouped: dict[tuple[str, str], list[float]] = {}
    for row in table.rows:
        try:
            value = float(row[value_index])
        except (TypeError, ValueError):
            continue
        grouped.setdefault((str(row[row_index]), str(row[column_index])), []).append(value)
    result: dict[str, dict[str, float]] = {}
    for (row_key, column_key), values in grouped.items():
        if operation == "sum":
            aggregate = sum(values)
        elif operation == "mean":
            aggregate = sum(values) / len(values)
        elif operation == "count":
            aggregate = float(len(values))
        elif operation == "min":
            aggregate = min(values)
        else:
            aggregate = max(values)
        result.setdefault(row_key, {})[column_key] = aggregate
    return result


def create_chart(
    context: SkillExecutionContext,
    path: str,
    labels: Sequence[str],
    values: Sequence[float],
    title: str,
    kind: str = "bar",
) -> FileInfo:
    """Create a validated PNG chart without accepting executable code."""
    if kind not in {"bar", "line"}:
        raise ToolError("chart kind must be bar or line")
    if not labels or len(labels) != len(values) or len(labels) > 1000:
        raise ToolError("chart labels and values must be non-empty and aligned")
    target = context._prepare_write(path)
    if target.suffix.lower() != ".png":
        raise ToolError("chart output must use .png")
    try:
        import matplotlib

        matplotlib.use("Agg")
        from matplotlib import pyplot as plt
    except ImportError as exc:
        raise ToolError("chart support requires matplotlib") from exc
    figure, axis = plt.subplots(figsize=(9, 5))
    if kind == "bar":
        axis.bar([str(label) for label in labels], [float(value) for value in values])
    else:
        axis.plot([str(label) for label in labels], [float(value) for value in values], marker="o")
    axis.set_title(str(title))
    axis.grid(axis="y", alpha=0.25)
    figure.tight_layout()
    figure.savefig(target, format="png", dpi=144)
    plt.close(figure)
    return FileInfo(target.relative_to(context.root).as_posix(), target.stat().st_size, "image/png")


class OfficeToolset:
    """Typed convenience facade used by model adapters and tests."""

    def __init__(self, context: SkillExecutionContext):
        self.context = context

    def list_files(self, directory: str = "input") -> list[FileInfo]:
        return list_files(self.context, directory)

    def read_file(self, path: str | Path) -> bytes:
        return read_file(self.context, path)

    def read_text(self, path: str | Path) -> str:
        return read_text(self.context, path)

    def write_file(self, path: str | Path, content: bytes | bytearray) -> FileInfo:
        return write_file(self.context, path, content)

    def write_text(self, path: str | Path, content: str) -> FileInfo:
        return write_text(self.context, path, content)

    def copy_file(self, source: str | Path, destination: str | Path) -> FileInfo:
        return copy_file(self.context, source, destination)

    def create_docx(self, path: str, title: str, paragraphs: Sequence[str]) -> FileInfo:
        return create_docx(self.context, path, title, paragraphs)

    def create_xlsx(self, path: str, columns: Sequence[str], rows: Sequence[Sequence[Any]], sheet_name: str = "Sheet1") -> FileInfo:
        return create_xlsx(self.context, path, columns, rows, sheet_name)

    def create_pptx(self, path: str, title: str, slides: Sequence[Mapping[str, Any]]) -> FileInfo:
        return create_pptx(self.context, path, title, slides)

    def create_pdf(self, path: str, title: str, paragraphs: Sequence[str]) -> FileInfo:
        return create_pdf(self.context, path, title, paragraphs)

    def read_docx(self, path: str | Path) -> DocumentData:
        return read_docx(self.context, path)

    def read_xlsx(self, path: str | Path) -> dict[str, TableData]:
        return read_xlsx(self.context, path)

    def read_pptx(self, path: str | Path) -> PresentationData:
        return read_pptx(self.context, path)

    def read_pdf(self, path: str | Path) -> PdfData:
        return read_pdf(self.context, path)

    def read_table(self, path: str | Path) -> TableData:
        return read_table(self.context, path)

    def profile_table(self, path: str | Path) -> dict[str, Any]:
        return profile_table(self.context, path)

    def aggregate_table(self, path: str, group_by: str, value_column: str, operation: str = "sum") -> dict[str, float]:
        return aggregate_table(self.context, path, group_by, value_column, operation)

    def pivot_table(self, path: str, row_column: str, column_column: str, value_column: str, operation: str = "sum") -> dict[str, dict[str, float]]:
        return pivot_table(self.context, path, row_column, column_column, value_column, operation)

    def create_chart(self, path: str, labels: Sequence[str], values: Sequence[float], title: str, kind: str = "bar") -> FileInfo:
        return create_chart(self.context, path, labels, values, title, kind)

    def publish_artifact(self, path: str | Path, title: str | None = None, mime_type: str | None = None) -> PublishedArtifact:
        return publish_artifact(self.context, path, title, mime_type)

    def run_skill_script(self, argv: Sequence[str], **kwargs: Any) -> ScriptResult:
        return run_skill_script(self.context, argv, **kwargs)


def build_tools(context: SkillExecutionContext) -> dict[str, Callable[..., Any]]:
    """Return the fixed typed tool map for one execution context."""
    toolset = OfficeToolset(context)
    return {
        "list_files": toolset.list_files,
        "read_file": toolset.read_file,
        "read_text": toolset.read_text,
        "write_file": toolset.write_file,
        "write_text": toolset.write_text,
        "copy_file": toolset.copy_file,
        "create_docx": toolset.create_docx,
        "create_xlsx": toolset.create_xlsx,
        "create_pptx": toolset.create_pptx,
        "create_pdf": toolset.create_pdf,
        "read_docx": toolset.read_docx,
        "read_xlsx": toolset.read_xlsx,
        "read_pptx": toolset.read_pptx,
        "read_pdf": toolset.read_pdf,
        "read_table": toolset.read_table,
        "profile_table": toolset.profile_table,
        "aggregate_table": toolset.aggregate_table,
        "pivot_table": toolset.pivot_table,
        "create_chart": toolset.create_chart,
        "publish_artifact": toolset.publish_artifact,
        "run_skill_script": toolset.run_skill_script,
    }


build_office_tools = build_tools
make_office_tools = build_tools
office_tools = build_tools


__all__ = [
    "ToolError",
    "PathBoundaryError",
    "SkillScriptDenied",
    "SkillScriptTimeout",
    "SkillOutputLimitExceeded",
    "SkillScriptError",
    "ScriptResult",
    "PublishedArtifact",
    "FileInfo",
    "TableData",
    "DocumentData",
    "PresentationData",
    "PdfData",
    "SkillExecutionContext",
    "run_skill_script",
    "list_files",
    "read_file",
    "read_text",
    "write_file",
    "write_text",
    "copy_file",
    "create_docx",
    "create_xlsx",
    "create_pptx",
    "create_pdf",
    "publish_artifact",
    "read_docx",
    "read_xlsx",
    "read_pptx",
    "read_pdf",
    "read_table",
    "profile_table",
    "aggregate_table",
    "pivot_table",
    "create_chart",
    "OfficeToolset",
    "build_tools",
    "build_office_tools",
    "make_office_tools",
    "office_tools",
]
